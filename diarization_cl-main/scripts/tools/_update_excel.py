# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook(r"c:\Users\EDY\Desktop\new\docs\会议ASR系统_完整实现方案.xlsx")

new_note = (
    "\n【实测更新 - 佩戴者实验】本数据集：zjl/ash戴116(顾客麦)，wzy戴118(店员麦)，"
    "每个文件=双人对话，麦克风戴在谁身上谁清晰，远处那个人分离不出。"
    "实测(2m)：116@2m=3标签{2,4,6}碎片；118@2m=1标签(远处顾客从未被独立分出)。"
    "FRCRN降噪对【找回远处说话人】无效(两副设备皆是)。"
    "结论：单通道仅保留近处说话人；远处说话人需要①近处另置麦克风 ②目标说话人提取(TSE)。"
)

ws = wb["总流程"]
for row in ws.iter_rows(min_row=2):
    if row[1].value == "B05":
        row[6].value = (row[6].value or "") + new_note

ws3 = wb["实施排期与验证"]
for row in ws3.iter_rows(min_row=2):
    if row[0].value == "M3":
        row[5].value = (row[5].value or "") + (
            "；实测2m两副麦克风远端均无法分离(116碎片化/118合一)，FRCRN降噪无效，需改采集或引入TSE。")

wb.save(r"c:\Users\EDY\Desktop\new\docs\会议ASR系统_完整实现方案.xlsx")
print("saved")
