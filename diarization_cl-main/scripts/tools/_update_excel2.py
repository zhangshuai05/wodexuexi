# -*- coding: utf-8 -*-
"""更新 Excel：阶段B 增加 多麦克风空间分离 路线；技术选型增行"""
import openpyxl
P = r"c:\Users\EDY\Desktop\new\docs\会议ASR系统_完整实现方案.xlsx"
wb = openpyxl.load_workbook(P)

# ---- Sheet1 总流程 ----
ws = wb["总流程"]
# 找到 B05 行，在其后插入 B06（多麦克风空间分离）
new_row = [
    "B. 说话人分离（核心）", "B06", "多麦克风空间分离（进阶方案）",
    "单通道保近弃远（实测确认），远处说话人需靠多麦克风在空间上筛选。原理＝DOA定向 + 波束成形增强：先用方向估计算法算出语音来向，再把阵列指向该方向、压制其他方向，把远处弱声从干扰中拎出来。解决118@2m<远处顾客未被分离>的根本途径。",
    "DOA：SRP-PHAT（人声更稳，首选）/ MUSIC / ESPRIT；波束成形：延迟求和(DS) / MVDR / GSC。可用麦克风：2~4麦阵列（间距拉大→角分辨率高）",
    "调研：①SRP-PHAT vs MUSIC 在强混响下的稳定性；②角分辨率与阵列孔径的关系；③多设备无相位同步时的【软融合】(先各自diarization再按内容合并) vs 严格波束成形",
    "低成本路径：先用现有【双视角数据】(同场景被两副设备同时录到,顾客+店员)做多设备融合验证,再决定是否购阵列硬件。物理规律：波束成形能增强【方向不同】的目标,不能凭空创造SNR；混响、源数未知、阵列孔径是三大约束。",
]
# 找 B05 所在行，在其后插入
insert_idx = None
for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
    if row[1].value == "B05":
        insert_idx = i + 1  # 在下一行插入
        break
if insert_idx:
    ws.insert_rows(insert_idx, 1)
    for c, v in enumerate(new_row, start=1):
        ws.cell(row=insert_idx, column=c, value=v)
    # 补样式（简单应用，忽略合并调整）
    for c in range(1, 8):
        cell = ws.cell(row=insert_idx, column=c)
        cell.font = openpyxl.styles.Font(name="微软雅黑", size=10)
        cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")

# B03 备注补充单通道局限指向 B06
for row in ws.iter_rows(min_row=2):
    if row[1].value == "B03":
        row[6].value = (row[6].value or "") + " 注：单通道仅保近处说话人，远处需多麦，见 B06。"

# ---- Sheet2 技术选型与调研 加行 ----
ws2 = wb["技术选型与调研"]
ws2.append([
    "空间分离", "DOA 方向估计", "SRP-PHAT（强混响更稳）", "MUSIC / ESPRIT",
    "numpy+scipy 自写（或 PyRoomAcoustics）",
    "本场景强混响下 SRP-PHAT vs MUSIC 实测对比", "SRP-PHAT 对人声鲁棒性更高；MUSIC 对源数/阵列模型敏感，游戏厅强混响易虚警。",
])
ws2.append([
    "空间分离", "波束成形提取", "延迟求和 / MVDR（自写）", "GSC / 自适应波束",
    "scipy.signal 实现，或 pyroomacoustics / torchaudio",
    "角分辨率 vs 阵列孔径（麦克风间距）", "阵列指向目标方向，压制其他方向。孔径越大角分辨率越高，代价是硬件/摆放空间。",
])
ws2.append([
    "空间分离", "多设备软融合", "各自diarization + 标签合并（现有双视角数据可直接试）", "严格时间对齐波束成形",
    "（现有：shuffle排列双通道）",
    "无相位同步时软融合的可靠性；时间轴不对齐如何匹配说话人",
    "现有数据的零成本验证：116=顾客声道,118=店员声道，先各自分离再按内容合并，看能否救回2m远处人。",
])

wb.save(P)
print("saved")
